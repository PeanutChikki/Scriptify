import re
import subprocess
import sys
from pathlib import Path

from docx import Document  # pip install python-docx

MODEL = "hf.co/AtomicChat/Qwen3.5-4B-GGUF:Q4_K_M"
INPUT_DOCX = Path.cwd() / "input.docx"  # always the current folder
MAX_ROUNDS = 10  # safety net so a confused model can't loop forever

# Placeholder styles recognised: {{name}}  {name}  [name]  <<name>>
# Change this regex if your template uses another style.
PLACEHOLDER_RE = re.compile(
    r"\{\{\s*([^{}]+?)\s*\}\}"
    r"|\{\s*([^{}]+?)\s*\}"
    r"|\[\s*([^\[\]]+?)\s*\]"
    r"|<<\s*([^<>]+?)\s*>>"
)

# Terminal junk / thinking-block cleanup (same as the reference script)
ANSI_RE = re.compile(r"\x1b\[[0-?]*[ -/]*[@-~]|\x1b\][^\x07]*\x07")
SPINNER_RE = re.compile(r"[⠋⠙⠹⠸⠼⠴⠦⠧⠇⠏]")
THINK_TAGS_RE = re.compile(r"<think>.*?</think>", re.DOTALL | re.IGNORECASE)
THINK_TEXT_RE = re.compile(r"Thinking\.\.\..*?\.\.\.done thinking\.", re.DOTALL)


def clean(text: str) -> str:
    text = ANSI_RE.sub("", text)
    text = SPINNER_RE.sub("", text)
    text = THINK_TAGS_RE.sub("", text)
    text = THINK_TEXT_RE.sub("", text)
    return text.strip()


def thought_rmv(text: str) -> str:
    tag = "</think>"
    index = text.find(tag)
    if index == -1:
        return text
    return text[index + len(tag):].strip()


def ask_model(prompt: str) -> str:
    # Prompt goes through stdin, so quotes/newlines are safe.
    result = subprocess.run(
        ["ollama", "run", MODEL, "--hidethinking", "--think=false"],
        input=prompt,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or "ollama exited with an error")
    return thought_rmv(clean(result.stdout))


# ---------------------------------------------------------------- docx reading
def iter_paragraph_texts(doc):
    """Yield the text of every paragraph: body, tables, headers and footers."""
    def from_container(container):
        for p in container.paragraphs:
            yield p.text
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from from_container(cell)

    yield from from_container(doc)
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer):
            yield from from_container(part)


def get_placeholders(path: Path) -> list[str]:
    """Return unique placeholder names in order of first appearance."""
    doc = Document(path)
    seen: dict[str, None] = {}
    for text in iter_paragraph_texts(doc):
        for match in PLACEHOLDER_RE.finditer(text):
            name = next(g for g in match.groups() if g)
            seen.setdefault(name.strip(), None)
    return list(seen)


# ---------------------------------------------------------------- AI checking
def build_check_prompt(indexed: dict[int, str], user_text: str) -> str:
    fields = "\n".join(f"{i}. {name}" for i, name in indexed.items())
    return f"""You are checking which form fields the user has supplied values for.

FIELDS (numbered):
{fields}

USER'S TEXT (everything they have written so far):
\"\"\"
{user_text}
\"\"\"

Decide which fields the user's text clearly gives a value for.
Never guess or invent a value. If a value is not clearly stated, that field is NOT found.

Respond with ONLY the numbers of the fields whose values were found, separated by commas.
Do not write anything else."""


def parse_reply(reply: str, valid: set[int]) -> set[int]:
    """Turn the model's reply (e.g. '1, 3, 4' or 'NONE') into a set of indexes."""
    if re.search(r"\d", reply):
        return {int(n) for n in re.findall(r"\d+", reply)} & valid
    if re.search(r"\bnone\b", reply, re.IGNORECASE):
        return set()
    raise ValueError("Model reply contained neither numbers nor NONE")


def find_found_indexes(indexed: dict[int, str], user_text: str) -> set[int]:
    prompt = build_check_prompt(indexed, user_text)
    valid = set(indexed)
    for attempt in range(3):  # small models sometimes ignore the format; retry
        reply = ask_model(prompt)
        try:
            return parse_reply(reply, valid)
        except ValueError:
            print(f"  (unexpected model reply, retrying {attempt + 1}/3...)")
    raise RuntimeError("Model failed to give a usable reply after 3 attempts")


# ----------------------------------------------------------------------- main
def main() -> None:
    if not INPUT_DOCX.exists():
        sys.exit(f"Could not find {INPUT_DOCX}")

    names = get_placeholders(INPUT_DOCX)
    if not names:
        sys.exit("No placeholders found in input.docx.")

    # Internal index: 1 -> first placeholder, 2 -> second, ...
    indexed = {i: name for i, name in enumerate(names, start=1)}

    print(f"Found {len(indexed)} placeholder(s) in {INPUT_DOCX.name}:")
    for i, name in indexed.items():
        print(f"  {i}. {name}")

    if len(sys.argv) > 1:
        user_text = " ".join(sys.argv[1:])
    else:
        user_text = input("\nEnter your prompt (give as many values as you can):\n> ").strip()
    if not user_text:
        sys.exit("No prompt given.")

    all_indexes = set(indexed)
    found: set[int] = set()
    try:
        for round_no in range(1, MAX_ROUNDS + 1):
            print(f"\nChecking values (round {round_no})...")
            found = find_found_indexes(indexed, user_text)
            missing = sorted(all_indexes - found)  # absent = everything not reported found

            print(f"AI reported found indexes: {sorted(found) or 'none'}")

            if not missing:
                print("\nAll values have been found.")
                break

            print("\nStill missing:")
            for i in missing:
                print(f"  {i}. {indexed[i]}")
            extra = input("\nPlease provide the missing values:\n> ").strip()
            if not extra:
                print("(nothing entered, asking again)")
            user_text += "\n" + extra
        else:
            print(f"\nStopped after {MAX_ROUNDS} rounds. Some values may still be missing.")
    except FileNotFoundError:
        sys.exit("Could not find 'ollama'. Is it installed and on your PATH?")
    except RuntimeError as err:
        sys.exit(f"Ollama error: {err}")

    # Save everything the user typed so the next step can extract/fill the values
    out = Path.cwd() / "user_input.txt"
    out.write_text(user_text, encoding="utf-8")
    print(f"\nCollected user input saved to: {out}")


if __name__ == "__main__":
    main()