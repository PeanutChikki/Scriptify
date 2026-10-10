import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

TXT_FILE = Path.cwd() / "output.txt"      # text to insert
TEMPLATE = Path.cwd() / "input.docx"      # your formatted document (never modified)
DOCX_FILE = Path.cwd() / "output.docx"    # result


def has_graphic(run) -> bool:
    """True if the run holds an image, shape, text box or embedded object."""
    return bool(run._r.xpath(".//w:drawing | .//w:pict | .//w:object"))


def text_runs(paragraph):
    """Runs that carry only text (so they are safe to overwrite)."""
    return [r for r in paragraph.runs if not has_graphic(r)]


def is_slot(paragraph) -> bool:
    """A paragraph can receive text unless it holds only an image/graphic."""
    if text_runs(paragraph):
        return True
    return not any(has_graphic(r) for r in paragraph.runs)


def set_text(paragraph, text: str) -> None:
    """Replace the paragraph's text, keeping its style and the first run's formatting.
    Runs containing images/shapes are left untouched."""
    runs = text_runs(paragraph)
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)


def insert_after(paragraph, text: str) -> Paragraph:
    """Add a new paragraph right after `paragraph`, copying its formatting."""
    new_p = OxmlElement("w:p")
    if paragraph._p.pPr is not None:
        ppr = deepcopy(paragraph._p.pPr)
        for sect in ppr.findall(qn("w:sectPr")):  # don't duplicate section breaks
            ppr.remove(sect)
        new_p.append(ppr)
    paragraph._p.addnext(new_p)

    new_par = Paragraph(new_p, paragraph._parent)
    run = new_par.add_run(text)

    src_runs = text_runs(paragraph)
    if src_runs and src_runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(src_runs[0]._r.rPr))
    return new_par


def txt_to_doc(
    txt_path: str | Path | None = None,
    template_path: str | Path | None = None,
    output_path: str | Path | None = None,
) -> Path:
    source_txt = Path(txt_path).resolve() if txt_path else TXT_FILE
    source_tmpl = Path(template_path).resolve() if template_path else TEMPLATE
    dest_docx = Path(output_path).resolve() if output_path else DOCX_FILE

    if not source_txt.exists():
        raise FileNotFoundError(f"Could not find text source: {source_txt}")
    if not source_tmpl.exists():
        raise FileNotFoundError(f"Could not find template file: {source_tmpl}")

    lines = source_txt.read_text(encoding="utf-8").splitlines()

    doc = Document(str(source_tmpl))

    # Only body paragraphs are touched. Tables, headers, footers, images,
    # page setup, styles and section breaks are never modified.
    slots = [p for p in doc.paragraphs if is_slot(p)]
    if not slots:
        # Fallback: if template has no paragraphs, create them
        for line in lines:
            doc.add_paragraph(line)
    else:
        # 1) Fill existing paragraphs in order
        for slot, line in zip(slots, lines):
            set_text(slot, line)

        # 2) Not enough paragraphs? Add the extra lines right after the last one used
        if len(lines) > len(slots):
            last = slots[-1]
            for line in lines[len(slots):]:
                last = insert_after(last, line)

        # 3) Too many paragraphs? Blank the leftovers (text only, images stay)
        elif len(lines) < len(slots):
            for slot in slots[len(lines):]:
                for r in text_runs(slot):
                    r.text = ""

    dest_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest_docx))
    print(f"Done. Saved to: {dest_docx}")
    return dest_docx
